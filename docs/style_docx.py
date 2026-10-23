from docx import Document
from docx.shared import Inches, Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path
import sys


def set_cell_shading(cell, fill='EDEDED'):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd')
        tcPr.append(shd)
    shd.set(qn('w:fill'), fill)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run('Page ')
    run.font.size = Pt(9)
    fldChar1 = OxmlElement('w:fldChar'); fldChar1.set(qn('w:fldCharType'), 'begin')
    instrText = OxmlElement('w:instrText'); instrText.set(qn('xml:space'), 'preserve'); instrText.text = ' PAGE '
    fldChar2 = OxmlElement('w:fldChar'); fldChar2.set(qn('w:fldCharType'), 'end')
    run._r.append(fldChar1); run._r.append(instrText); run._r.append(fldChar2)


def style_docx(src, dst, short_title):
    doc = Document(src)
    sec = doc.sections[0]
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    sec.top_margin = Cm(2.0)
    sec.bottom_margin = Cm(2.0)
    sec.left_margin = Cm(2.1)
    sec.right_margin = Cm(2.1)
    sec.header_distance = Cm(0.8)
    sec.footer_distance = Cm(0.8)

    styles = doc.styles
    normal = styles['Normal']
    normal.font.name = 'Times New Roman'
    normal._element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.08
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    for sty, size, bold in [('Title', 16, True), ('Heading 1', 13, True), ('Heading 2', 11.5, True), ('Heading 3', 10.5, True)]:
        if sty in styles:
            s = styles[sty]
            s.font.name = 'Times New Roman'
            s._element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
            s.font.size = Pt(size)
            s.font.bold = bold
            s.paragraph_format.keep_with_next = True
            s.paragraph_format.space_before = Pt(8 if sty != 'Title' else 0)
            s.paragraph_format.space_after = Pt(4)

    if 'Subtitle' in styles:
        s = styles['Subtitle']; s.font.name='Times New Roman'; s.font.size=Pt(10.5)
    if 'Caption' in styles:
        s=styles['Caption']; s.font.name='Times New Roman'; s.font.size=Pt(9); s.font.italic=True
        s.paragraph_format.space_before=Pt(3); s.paragraph_format.space_after=Pt(6)

    # First title/author paragraphs center when present.
    for p in doc.paragraphs[:10]:
        if p.style.name in ('Title','Subtitle','Author'):
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Keep figure-containing paragraphs together and center.
    for p in doc.paragraphs:
        has_drawing = bool(p._p.xpath('.//w:drawing'))
        if has_drawing:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.keep_together = True
            p.paragraph_format.space_before = Pt(5)
            p.paragraph_format.space_after = Pt(3)
        # code/preformatted
        if p.style.name in ('Source Code','Verbatim','Code Block'):
            for r in p.runs:
                r.font.name='Courier New'; r.font.size=Pt(8.5)

    # Tables
    for table in doc.tables:
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = True
        for ri,row in enumerate(table.rows):
            for cell in row.cells:
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                for p in cell.paragraphs:
                    p.paragraph_format.space_after = Pt(1)
                    p.paragraph_format.line_spacing = 1.0
                    for r in p.runs:
                        r.font.name='Times New Roman'; r.font.size=Pt(8.5)
                        if ri == 0: r.font.bold=True
                if ri == 0:
                    set_cell_shading(cell, 'E8E8E8')
        # repeat header
        if table.rows:
            trPr = table.rows[0]._tr.get_or_add_trPr()
            tblHeader = OxmlElement('w:tblHeader'); tblHeader.set(qn('w:val'),'true'); trPr.append(tblHeader)

    # Avoid stretched justification in metadata and reference-list paragraphs.
    ref_mode = False
    author_mode = False
    for p in doc.paragraphs:
        txt = p.text.strip()
        if txt == 'Author information':
            author_mode = True
        elif txt == 'References':
            ref_mode = True
            author_mode = False
        elif p.style.name.startswith('Heading') and txt not in ('Author information','References'):
            author_mode = False
        if ref_mode or author_mode:
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        if ref_mode and not txt == 'References':
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.line_spacing = 1.0
            for r in p.runs:
                r.font.size = Pt(9)

    # Header/footer
    header = sec.header
    hp = header.paragraphs[0]
    hp.text = short_title
    hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for r in hp.runs:
        r.font.name='Times New Roman'; r.font.size=Pt(8); r.font.italic=True
    footer = sec.footer
    fp = footer.paragraphs[0]
    add_page_number(fp)
    for r in fp.runs:
        r.font.name='Times New Roman'; r.font.size=Pt(8)

    # Core properties
    doc.core_properties.author = 'Emirhan Yıldırım'
    doc.core_properties.title = short_title
    doc.core_properties.subject = 'Invariant geometry; zonotopes; quotient invariants; reproducibility'
    doc.save(dst)

if __name__ == '__main__':
    style_docx(sys.argv[1], sys.argv[2], sys.argv[3])
